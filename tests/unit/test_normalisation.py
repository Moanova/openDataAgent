# -*- coding: utf-8 -*-
"""Tests de la normalisation B7 (T10.1) — TU-124..TU-133, TI-27.

Couvre `src/core/normalisation.py` : détection de conformité CSV,
normalisation CSV (encodage, séparateur), conversion tableurs
(xlsx/xls/ods) en un CSV par feuille, copie des bruts non normalisables,
exclusion du manifeste, re-lancement sans doublon, simulation, bilan
normalisé, commande CLI `normaliser` (une exécution, toutes, --json).
"""
import csv
import json
import zipfile
from pathlib import Path

import pytest
from typer.testing import CliRunner

from src.core.normalisation import (
    RACINE_BRUT,
    RACINE_TRAITE,
    _est_csv_conforme,
    normaliser_execution,
    normaliser_toutes_executions,
)
from src.main import app

runner = CliRunner()

EXEC_ID = "20261023_100000_a1b2c3"


# ---------------------------------------------------------------------------
# Fixtures : fabrique des classeurs et un ODS minimal valides
# ---------------------------------------------------------------------------

def _fabriquer_xlsx(chemin: Path, feuilles: dict) -> None:
    from openpyxl import Workbook

    classeur = Workbook()
    premiere = True
    for nom, lignes in feuilles.items():
        feuille = classeur.active if premiere else classeur.create_sheet()
        premiere = False
        feuille.title = nom
        for ligne in lignes:
            feuille.append(ligne)
    classeur.save(chemin)


def _fabriquer_xls(chemin: Path, feuilles: dict) -> None:
    import xlwt

    classeur = xlwt.Workbook()
    for nom, lignes in feuilles.items():
        feuille = classeur.add_sheet(nom)
        for r, ligne in enumerate(lignes):
            for c, valeur in enumerate(ligne):
                feuille.write(r, c, valeur)
    classeur.save(chemin)


def _fabriquer_ods(chemin: Path, feuilles: dict) -> None:
    def cellule(valeur):
        if isinstance(valeur, (int, float)):
            return (f'<table:table-cell office:value-type="float" office:value="{valeur}">'
                    f"<text:p>{valeur}</text:p></table:table-cell>")
        return f'<table:table-cell office:value-type="string"><text:p>{valeur}</text:p></table:table-cell>'

    tables = []
    for nom, lignes in feuilles.items():
        rangees = "".join(
            "<table:table-row>" + "".join(cellule(v) for v in ligne) + "</table:table-row>"
            for ligne in lignes
        )
        tables.append(f'<table:table table:name="{nom}">{rangees}</table:table>')
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content '
        'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
        'office:version="1.2"><office:body><office:spreadsheet>'
        + "".join(tables)
        + "</office:spreadsheet></office:body></office:document-content>"
    )
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2">'
        '<manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/>'
        '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
        "</manifest:manifest>"
    )
    with zipfile.ZipFile(chemin, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet", compress_type=zipfile.ZIP_STORED)
        z.writestr("content.xml", content)
        z.writestr("META-INF/manifest.xml", manifest)


@pytest.fixture
def dossier_execution(tmp_path):
    """data/raw/<EXEC_ID>/ avec un panel de fichiers + manifeste."""
    dossier = tmp_path / "data" / "raw" / EXEC_ID
    dossier.mkdir(parents=True)
    (dossier / "71ff1335_donnees.csv").write_bytes("nom,age\nalice,30\n".encode("utf-8"))
    (dossier / "abc123_budget.csv").write_bytes("nom;ville\nalicé;Marseille\nzoé;Lyon\n".encode("latin-1"))
    (dossier / "jkl012_archive.zip").write_bytes(b"PK\x03\x04faux")
    (dossier / "manifeste_acquisitions.json").write_text("{}", encoding="utf-8")
    return tmp_path


def _bilan(bilan, action, fichier):
    for f in bilan["donnees"]["fichiers"]:
        if f["fichier"] == fichier:
            return f
    raise AssertionError(f"fichier {fichier} absent du bilan (action={action})")


# ---------------------------------------------------------------------------
# TU-124 — Détection de conformité CSV
# ---------------------------------------------------------------------------

def test_tu124_csv_conforme_detecte(tmp_path):
    # TU-124 : un CSV UTF-8 virgule sans BOM est détecté conforme.
    conforme = tmp_path / "ok.csv"
    conforme.write_bytes("nom,age\nalice,30\n".encode("utf-8"))
    assert _est_csv_conforme(conforme) is True

    latin1 = tmp_path / "latin1.csv"
    latin1.write_bytes("nom;ville\nalicé;Lyon\n".encode("latin-1"))
    assert _est_csv_conforme(latin1) is False

    bom = tmp_path / "bom.csv"
    bom.write_bytes(b"\xef\xbb\xbf" + "nom,age\n".encode("utf-8"))
    assert _est_csv_conforme(bom) is False

    point_virgule = tmp_path / "pv.csv"
    point_virgule.write_bytes("nom;age\nalice;30\n".encode("utf-8"))
    assert _est_csv_conforme(point_virgule) is False

    vide = tmp_path / "vide.csv"
    vide.write_bytes(b"")
    assert _est_csv_conforme(vide) is True


# ---------------------------------------------------------------------------
# TU-125 — Normalisation CSV (encodage + séparateur)
# ---------------------------------------------------------------------------

def test_tu125_csv_latin1_point_virgule_normalise(dossier_execution):
    # TU-125 : un CSV latin-1 point-virgule est ré-écrit en UTF-8 virgule.
    bilan = normaliser_execution(
        EXEC_ID,
        racine_brut=dossier_execution / "data" / "raw",
        racine_traite=dossier_execution / "data" / "processed",
    )
    assert bilan["statut"] == "succes"
    resultat = _bilan(bilan, "normalise", "abc123_budget.csv")
    assert resultat["action"] == "normalise"

    sortie = dossier_execution / "data" / "processed" / EXEC_ID / "abc123_budget.csv"
    contenu = sortie.read_text(encoding="utf-8")
    assert contenu == "nom,ville\nalicé,Marseille\nzoé,Lyon\n"
    assert not sortie.read_bytes().startswith(b"\xef\xbb\xbf")


def test_tu125b_csv_bom_normalise(tmp_path):
    # TU-125 (suite) : un BOM UTF-8 est retiré à la ré-écriture.
    dossier = tmp_path / "data" / "raw" / EXEC_ID
    dossier.mkdir(parents=True)
    (dossier / "bom.csv").write_bytes(b"\xef\xbb\xbf" + "a,b\n1,2\n".encode("utf-8"))
    bilan = normaliser_execution(
        EXEC_ID,
        racine_brut=tmp_path / "data" / "raw",
        racine_traite=tmp_path / "data" / "processed",
    )
    sortie = tmp_path / "data" / "processed" / EXEC_ID / "bom.csv"
    assert sortie.read_bytes() == "a,b\n1,2\n".encode("utf-8")
    assert bilan["statut"] == "succes"


# ---------------------------------------------------------------------------
# TU-126 — CSV déjà conforme : copié tel quel (voie 1, pas de retraitement)
# ---------------------------------------------------------------------------

def test_tu126_csv_conforme_copie_tel_quel(dossier_execution):
    # TU-126 : le CSV pivot de la voie 1 (déjà UTF-8 virgule) est copié
    # tel quel, sans ré-écriture.
    source = dossier_execution / "data" / "raw" / EXEC_ID / "71ff1335_donnees.csv"
    bilan = normaliser_execution(
        EXEC_ID,
        racine_brut=dossier_execution / "data" / "raw",
        racine_traite=dossier_execution / "data" / "processed",
    )
    resultat = _bilan(bilan, "conforme", "71ff1335_donnees.csv")
    assert resultat["action"] == "conforme"
    sortie = dossier_execution / "data" / "processed" / EXEC_ID / "71ff1335_donnees.csv"
    assert sortie.read_bytes() == source.read_bytes()


# ---------------------------------------------------------------------------
# TU-127 — Conversion tableurs (xlsx, xls, ods) : un CSV par feuille
# ---------------------------------------------------------------------------

FEUILLES = {
    "Feuille1": [["nom", "age"], ["alice", 30]],
    "Stats": [["total", 130]],
}
FEUILLES_CSV = {
    "Feuille1": [["nom", "age"], ["alice", "30"]],
    "Stats": [["total", "130"]],
}
# xlrd et odfpy lisent les nombres comme des flottants ("30.0", "130.0") ;
# openpyxl conserve l'entier. Chaque lecteur a sa représentation attendue.
FEUILLES_CSV_XLRD_ODS = {
    "Feuille1": [["nom", "age"], ["alice", "30.0"]],
    "Stats": [["total", "130.0"]],
}


def _lire_csv(chemin: Path):
    with chemin.open(encoding="utf-8", newline="") as flux:
        return list(csv.reader(flux))


@pytest.mark.parametrize("extension,fabriquer", [
    (".xlsx", _fabriquer_xlsx),
    (".xls", _fabriquer_xls),
    (".ods", _fabriquer_ods),
])
def test_tu127_tableur_convertis_par_feuille(tmp_path, extension, fabriquer):
    # TU-127 : xlsx, xls et ods produisent un CSV par feuille, nommé
    # <fichier>__<feuille>.csv, en UTF-8 virgule.
    dossier = tmp_path / "data" / "raw" / EXEC_ID
    dossier.mkdir(parents=True)
    fabriquer(dossier / f"def456_classeur{extension}", FEUILLES)

    bilan = normaliser_execution(
        EXEC_ID,
        racine_brut=tmp_path / "data" / "raw",
        racine_traite=tmp_path / "data" / "processed",
    )
    assert bilan["statut"] == "succes"
    resultat = _bilan(bilan, "convertis", f"def456_classeur{extension}")
    assert resultat["action"] == "convertis"
    produits = {p["fichier"]: p for p in resultat["detail"]}
    assert set(produits) == {
        "def456_classeur__Feuille1.csv",
        "def456_classeur__Stats.csv",
    }
    destination = tmp_path / "data" / "processed" / EXEC_ID
    attendu = FEUILLES_CSV if extension == ".xlsx" else FEUILLES_CSV_XLRD_ODS
    assert _lire_csv(destination / "def456_classeur__Feuille1.csv") == attendu["Feuille1"]
    assert _lire_csv(destination / "def456_classeur__Stats.csv") == attendu["Stats"]
    # Le classeur source n'est pas copié dans processed/
    assert not (destination / f"def456_classeur{extension}").exists()


def test_tu127b_nom_feuille_avec_espaces_nettoye(tmp_path):
    # TU-127 (suite) : les espaces/accents du nom de feuille sont nettoyés
    # dans le nom de fichier produit.
    dossier = tmp_path / "data" / "raw" / EXEC_ID
    dossier.mkdir(parents=True)
    _fabriquer_xlsx(dossier / "classeur.xlsx", {"Données 2": [["x"]]})
    normaliser_execution(
        EXEC_ID,
        racine_brut=tmp_path / "data" / "raw",
        racine_traite=tmp_path / "data" / "processed",
    )
    assert (tmp_path / "data" / "processed" / EXEC_ID / "classeur__Donn-es-2.csv").is_file()


# ---------------------------------------------------------------------------
# TU-128 — Bruts non normalisables copiés tels quels
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("nom,contenu", [
    ("archive.zip", b"PK\x03\x04faux"),
    ("donnees.geojson", b'{"type": "FeatureCollection", "features": []}'),
    ("notes.tsv", b"a\tb\n1\t2\n"),
    ("rapport.pdf", b"%PDF-1.4 fake"),
    ("inconnu.xyz", b"\x00\x01\x02"),
])
def test_tu128_brut_non_normalisable_copie_tel_quel(dossier_execution, nom, contenu):
    # TU-128 : zip, geojson, tsv, pdf, extension inconnue : copiés tels quels.
    dossier = dossier_execution / "data" / "raw" / EXEC_ID
    dossier.joinpath(nom).write_bytes(contenu)
    bilan = normaliser_execution(
        EXEC_ID,
        racine_brut=dossier_execution / "data" / "raw",
        racine_traite=dossier_execution / "data" / "processed",
    )
    resultat = _bilan(bilan, "copie", nom)
    assert resultat["action"] == "copie"
    sortie = dossier_execution / "data" / "processed" / EXEC_ID / nom
    assert sortie.read_bytes() == contenu


# ---------------------------------------------------------------------------
# TU-129 — processed/ ne contient que des fichiers finaux (manifeste exclu)
# ---------------------------------------------------------------------------

def test_tu129_manifeste_exclu_de_processed(dossier_execution):
    # TU-129 : le manifeste d'acquisition n'est jamais copié dans processed/.
    normaliser_execution(
        EXEC_ID,
        racine_brut=dossier_execution / "data" / "raw",
        racine_traite=dossier_execution / "data" / "processed",
    )
    destination = dossier_execution / "data" / "processed" / EXEC_ID
    assert not (destination / "manifeste_acquisitions.json").exists()
    assert all(f.suffix in (".csv", ".zip") for f in destination.iterdir())


# ---------------------------------------------------------------------------
# TU-130 — Re-lancement sans doublon (jamais d'écrasement)
# ---------------------------------------------------------------------------

def test_tu130_relancement_sans_doublon(dossier_execution):
    # TU-130 : un second passage ne duplique ni n'écrase aucun fichier ;
    # chaque fichier est marqué 'ignore'.
    brut = dossier_execution / "data" / "raw"
    traite = dossier_execution / "data" / "processed"
    _fabriquer_xlsx(brut / EXEC_ID / "def456_classeur.xlsx", FEUILLES)

    premier = normaliser_execution(EXEC_ID, racine_brut=brut, racine_traite=traite)
    noms_avant = sorted(f.name for f in (traite / EXEC_ID).iterdir())
    empreintes_avant = {f.name: f.read_bytes() for f in (traite / EXEC_ID).iterdir()}

    second = normaliser_execution(EXEC_ID, racine_brut=brut, racine_traite=traite)
    noms_apres = sorted(f.name for f in (traite / EXEC_ID).iterdir())

    assert noms_avant == noms_apres
    for fichier in second["donnees"]["fichiers"]:
        assert fichier["action"] == "ignore"
    for nom, contenu in empreintes_avant.items():
        assert (traite / EXEC_ID / nom).read_bytes() == contenu


# ---------------------------------------------------------------------------
# TU-131 — Simulation : bilan fidèle, rien d'écrit
# ---------------------------------------------------------------------------

def test_tu131_simulation_necrit_rien(dossier_execution):
    # TU-131 : --simulation produit le même bilan qu'un passage réel mais
    # n'écrit aucun fichier sous data/processed/.
    brut = dossier_execution / "data" / "raw"
    traite = dossier_execution / "data" / "processed"
    bilan_sim = normaliser_execution(EXEC_ID, racine_brut=brut, racine_traite=traite, simulation=True)
    assert bilan_sim["statut"] == "succes"
    assert bilan_sim["donnees"]["simulation"] is True
    assert not traite.exists()
    actions_sim = {f["fichier"]: f["action"] for f in bilan_sim["donnees"]["fichiers"]}
    assert actions_sim["71ff1335_donnees.csv"] == "conforme"
    assert actions_sim["abc123_budget.csv"] == "normalise"
    assert actions_sim["jkl012_archive.zip"] == "copie"


# ---------------------------------------------------------------------------
# TU-132 — Bilan normalisé et exécution inconnue
# ---------------------------------------------------------------------------

def test_tu132_execution_inconnue_erreur_parametres(tmp_path):
    # TU-132 : une exécution sans dossier data/raw/<id>/ retourne
    # erreur_parametres (jamais d'exception).
    bilan = normaliser_execution(
        "20200101_000000_xxxxxx",
        racine_brut=tmp_path / "data" / "raw",
        racine_traite=tmp_path / "data" / "processed",
    )
    assert bilan["statut"] == "erreur_parametres"
    assert "20200101_000000_xxxxxx" in bilan["message"]


def test_tu132b_execution_sans_fichier_succes_vide(tmp_path):
    # TU-132 (suite) : dossier présent mais sans fichier (manifeste seul)
    # => succes, aucun fichier traité.
    dossier = tmp_path / "data" / "raw" / EXEC_ID
    dossier.mkdir(parents=True)
    (dossier / "manifeste_acquisitions.json").write_text("{}", encoding="utf-8")
    bilan = normaliser_execution(
        EXEC_ID,
        racine_brut=tmp_path / "data" / "raw",
        racine_traite=tmp_path / "data" / "processed",
    )
    assert bilan["statut"] == "succes"
    assert bilan["donnees"]["fichiers"] == []


# ---------------------------------------------------------------------------
# TU-133 — normaliser_toutes_executions : agrégation + exclusion courante
# ---------------------------------------------------------------------------

def test_tu133_toutes_executions_agregees(dossier_execution):
    # TU-133 : sans argument, toutes les exécutions de data/raw/ sont
    # normalisées ; l'exécution courante passée en paramètre est exclue.
    brut = dossier_execution / "data" / "raw"
    traite = dossier_execution / "data" / "processed"
    autre = brut / "20261022_090000_b2c3d4"
    autre.mkdir()
    (autre / "x.csv").write_bytes("a,b\n1,2\n".encode("utf-8"))

    bilan = normaliser_toutes_executions(
        racine_brut=brut,
        racine_traite=traite,
        id_execution_courante=EXEC_ID,
    )
    assert bilan["statut"] == "succes"
    assert bilan["donnees"]["executions"] == ["20261022_090000_b2c3d4"]
    assert (traite / "20261022_090000_b2c3d4" / "x.csv").is_file()
    assert not (traite / EXEC_ID).exists()


# ---------------------------------------------------------------------------
# TI-27 — Commande CLI `normaliser` (exécution, toutes, --json, simulation)
# ---------------------------------------------------------------------------

def test_ti27_cli_normaliser_cycle_complet(tmp_path, monkeypatch):
    # TI-27 : la commande `normaliser` enchaîne simulation (rien d'écrit),
    # exécution réelle, re-lancement (tout ignoré) et --json sérialisable.
    # L'option globale --silencieux est passée en argument CLI : la sortie
    # --json reste parsable même quand des tests précédents ont activé les
    # logs console (le callback Typer réinitialise _SILENCIEUX à chaque
    # invocation, un monkeypatch sur la variable globale serait sans effet).
    monkeypatch.chdir(tmp_path)
    dossier = tmp_path / "data" / "raw" / EXEC_ID
    dossier.mkdir(parents=True)
    (dossier / "71ff1335_donnees.csv").write_bytes("nom,age\nalice,30\n".encode("utf-8"))
    (dossier / "abc123_budget.csv").write_bytes("nom;ville\nalicé;Lyon\n".encode("latin-1"))
    _fabriquer_xlsx(dossier / "def456_classeur.xlsx", FEUILLES)

    # Simulation : rien d'écrit
    resultat = runner.invoke(app, ["--silencieux", "normaliser", EXEC_ID, "--simulation"])
    assert resultat.exit_code == 0
    assert "simulation : True" in resultat.output
    assert not (tmp_path / "data" / "processed").exists()

    # Réel
    resultat = runner.invoke(app, ["--silencieux", "normaliser", EXEC_ID])
    assert resultat.exit_code == 0
    assert "normalise" in resultat.output
    assert "convertis" in resultat.output
    destination = tmp_path / "data" / "processed" / EXEC_ID
    assert (destination / "abc123_budget.csv").read_text(encoding="utf-8") == "nom,ville\nalicé,Lyon\n"
    assert (destination / "def456_classeur__Feuille1.csv").is_file()

    # Re-lancement : tout ignoré, aucun doublon
    resultat = runner.invoke(app, ["--silencieux", "normaliser", EXEC_ID, "--json"])
    assert resultat.exit_code == 0
    bilan = json.loads(resultat.output)
    assert bilan["statut"] == "succes"
    assert all(f["action"] == "ignore" for f in bilan["donnees"]["fichiers"])
    assert len(list(destination.iterdir())) == 4  # 2 CSV + 2 feuilles produites

    # Sans argument : toutes les exécutions
    resultat = runner.invoke(app, ["--silencieux", "normaliser", "--json"])
    assert resultat.exit_code == 0
    bilan = json.loads(resultat.output)
    assert bilan["statut"] == "succes"
    assert EXEC_ID in bilan["donnees"]["executions"]

    # Exécution inconnue : erreur propre, code de sortie 1
    resultat = runner.invoke(app, ["--silencieux", "normaliser", "20200101_000000_xxxxxx"])
    assert resultat.exit_code == 1
    assert "erreur_parametres" in resultat.output
